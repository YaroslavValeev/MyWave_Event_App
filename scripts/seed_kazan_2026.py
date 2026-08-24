"""Seed Kazan 2026 Championship event from local Excel/PDF sources.

Privacy:
- Phones are imported onto Participant + User for phone-login MVP.
- Medical Drive URLs may be stored internally (medical_cert_url) but are NEVER logged
  and NEVER exposed via public ParticipantOut (only has_medical_cert bool).
- Phone numbers are never printed to stdout.
- Source files stay outside git; PDFs are copied into data/documents/.
"""

from __future__ import annotations

import argparse
import hashlib
import re
import shutil
import sys
from datetime import date, datetime, time, timezone
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import delete, select

ROOT = Path(__file__).resolve().parents[1]
API_DIR = ROOT / "services" / "api"
sys.path.insert(0, str(API_DIR))

from app.config import get_settings  # noqa: E402
from app.db import get_session_factory, init_db  # noqa: E402
from app.models import Category, Document, Event, Official, Participant, TrainingSlot, User  # noqa: E402
from app.models.event import EventStatus  # noqa: E402
from app.services.phone_utils import normalize_phone  # noqa: E402

DEFAULT_LEGACY_REG = Path(
    r"c:\Users\X230\Downloads\Telegram Desktop\Регистрация_ЧР_ПР_2026_по_категориям.xlsx"
)
DEFAULT_FORM_REG = Path(
    r"c:\Users\X230\Downloads\Предварительная регистрация ЧР и ПР 2026 - доски (Ответы) (1).xlsx"
)
DEFAULT_REG = DEFAULT_FORM_REG
TRAINING_XLSX = Path(r"c:\Users\X230\Downloads\2026 - Тренировки ПР ЧР.xlsx")
DESKTOP_BULLETIN = Path(
    r"c:\Users\X230\Desktop\ЧР_Казань\Бюллетень_Первенство_Чемпионат_России_2026_Wakeboard_Wakesurf_Wakeskim (2).pdf"
)
LEGACY_BULLETIN = Path(
    r"c:\Users\X230\Downloads\Telegram Desktop\Бюллетень_Первенство_Чемпионат_России_2026_Wakeboard_Wakesurf_Wakeskim (2).pdf"
)
JUDGE_PROTOCOL = Path(
    r"c:\Users\X230\Desktop\ЧР_Казань\32_протокол_КС_ЧР_Казань_доска_короткая,_длинная_260714_172154.pdf"
)

DEFAULT_DOCS = [
    (
        DESKTOP_BULLETIN if DESKTOP_BULLETIN.exists() else LEGACY_BULLETIN,
        "bulletin",
        "ru",
        "Бюллетень ЧР/ПР 2026 Wakeboard / Wakesurf / Wakeskim (обновлённый)",
    ),
    (
        JUDGE_PROTOCOL,
        "protocol",
        "ru",
        "Протокол №32-КС: судейский корпус (катер — доска длинная/короткая)",
    ),
    (
        TRAINING_XLSX,
        "schedule",
        "ru",
        "Расписание тренировок ПР/ЧР 2026 (Excel, актуальное)",
    ),
    (
        Path(r"c:\Users\X230\Downloads\2026-IWWF-Wakeboard-Boat-Rules-FINAL.pdf"),
        "rules",
        "en",
        "IWWF 2026 Wakeboard Boat Rules",
    ),
    (
        Path(r"c:\Users\X230\Downloads\2026-OFFICIAL-IWWF-WAKESURF-RULES-FINAL.pdf"),
        "rules",
        "en",
        "IWWF 2026 Official Wakesurf Rules",
    ),
]

# Из протокола №32-КС (заседание 13.07.2026): состав на Казань Wakesurf/Wakeskim
OFFICIALS = [
    (1, "Филиппов А. В.", "Главный судья", "г. Санкт-Петербург", "ВК, МК", "judge"),
    (2, "Живаев С. В.", "Главный секретарь", "Самарская область", "ВК", "judge"),
    (3, "Савичев А. А.", "Судья арбитр в катере", "г. Санкт-Петербург", "3К", "judge"),
    (4, "Козлов Д. А.", "Судья арбитр в катере", "г. Москва", "2К", "judge"),
    (5, "Живаев В. С.", "Судья арбитр в катере", "Самарская область", "2К", "judge"),
    (6, "Корольков С. И.", "Судья арбитр в катере", "г. Москва", "3К", "judge"),
    (7, "Симонов Е. В.", "Судья арбитр в катере", "г. Москва", "3К", "judge"),
    (8, "Шереметьев Н. К.", "Судья арбитр в катере", "г. Москва", "3К", "judge"),
    (9, "Пиоттух В. С.", "Судья водитель катера", "г. Москва", "3К", "judge"),
    (
        10,
        "По назначению ФВСРТ",
        "Судья при участниках",
        "Республика Татарстан",
        "3К",
        None,
    ),
    (11, "Валеев Я.", "Судья комментатор", "Республика Татарстан", None, "commentator"),
]

SHEET_DISCIPLINE = {
    "Вейкборд-катер": "Wakeboard (boat)",
    "Wakesurf (доска длинная)": "Wakesurf",
    "Wakeskim (доска короткая)": "Wakeskim",
}

FORM_DISCIPLINE_MAP = {
    "вейкборд - катер": "Wakeboard (boat)",
    "вейкборд-катер": "Wakeboard (boat)",
    "катер - доска длинная (wakesurf)": "Wakesurf",
    "катер - доска короткая (wakeskim)": "Wakeskim",
}

HEADER_ALIASES = {
    "fio": {"фио", "ф.и.о.", "русск", "fullname"},
    "gender": {"пол"},
    "category": {"категория"},
    "region": {"регион"},
    "birth": {"дата рождения"},
    "phone": {"телефон", "phone", "мобил", "контактный"},
    "discipline": {"дисциплина"},
    "medical": {"медицин", "справка", "скан-копия"},
}


def _slugify(value: str) -> str:
    digest = hashlib.sha1(value.encode("utf-8")).hexdigest()[:10]
    ascii_bits = re.sub(r"[^a-z0-9]+", "-", value.lower())
    ascii_bits = re.sub(r"-+", "-", ascii_bits).strip("-")[:40]
    return f"{ascii_bits}-{digest}" if ascii_bits else digest


def _normalize_fio(value: str) -> str:
    return re.sub(r"\s+", " ", value.strip().casefold())


def _fio_keys(value: str) -> list[str]:
    """Exact key + surname+name compact key for cross-sheet matching."""
    norm = _normalize_fio(value)
    keys = [norm]
    parts = [p for p in re.split(r"[\s.]+", norm) if p]
    if len(parts) >= 2:
        keys.append(f"{parts[0]} {parts[1]}")
        keys.append(f"{parts[0]}{parts[1][0]}")
    return keys


def _parse_headers(row: list) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for idx, cell in enumerate(row):
        if cell is None:
            continue
        label = str(cell).strip().lower()
        for key, aliases in HEADER_ALIASES.items():
            if label in aliases or any(a in label for a in aliases):
                mapping[key] = idx
    return mapping


def _is_section_row(row: list) -> bool:
    first = row[0] if row else None
    if first is None:
        return False
    text = str(first).strip()
    if not text:
        return False
    # Section titles look like "1. Девочки U14" without FIO in col2
    second = row[1] if len(row) > 1 else None
    return second is None or str(second).strip() == ""


def _birth_year(value) -> int | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.year
    text = str(value)
    m = re.search(r"(19|20)\d{2}", text)
    return int(m.group(0)) if m else None


def ensure_organizer(db) -> User:
    for email, role, name, phone in [
        ("organizer@example.com", "organizer", "Организатор ЧР/ПР 2026", None),
        ("y.valeev@gmail.com", "platform_admin", "Валеев Ярослав Радионович", "+79160117179"),
    ]:
        user = db.scalar(select(User).where(User.email == email))
        if user is None:
            db.add(
                User(
                    email=email,
                    role=role,
                    status="active",
                    display_name=name,
                    requested_role=role,
                    phone=phone,
                )
            )
            db.flush()
        else:
            user.role = role
            user.status = "active"
            user.display_name = name
            user.requested_role = role
            if phone:
                # Owner phone must win over any colliding seed user.
                collision = db.scalar(
                    select(User).where(User.phone == phone, User.id != user.id)
                )
                if collision is not None:
                    collision.phone = None
                    db.add(collision)
                user.phone = phone
            db.add(user)
            db.flush()
    return db.scalar(select(User).where(User.email == "organizer@example.com"))


def upsert_event(db) -> Event:
    slug = "chr-pr-kazan-2026"
    event = db.scalar(select(Event).where(Event.slug == slug))
    description = (
        "Чемпионат и Первенство России 2026 по воднолыжному спорту: "
        "вейкборд–катер, катер–доска длинная (Wakesurf), катер–доска короткая (Wakeskim). "
        "Период: 10–17 августа 2026, г. Казань, Республика Татарстан. "
        "Основная акватория: озеро Кабан, ул. Хади Такташа, 35. "
        "Официальные тренировки Wakesurf/Wakeskim: ул. Торфяная, 83. "
        "Расписание: 11.08 запасной тренировочный; 12.08 тренировки; "
        "13.08 квалификация; 14.08 полуфиналы; 15.08 финалы; 16.08 резерв. "
        "Судейство — по правилам IWWF. Регистрация закрывается 11.08.2026 00:00. "
        "Стартовый взнос 6500 ₽ (несовершеннолетние участники ПР — без стартового взноса)."
    )
    payload = dict(
        title="Чемпионат и Первенство России 2026 — Wakeboard / Wakesurf / Wakeskim",
        description=description,
        city="Казань",
        location="Республика Татарстан, г. Казань, озеро Кабан, ул. Хади Такташа, 35",
        venue="Оз. Кабан (Хади Такташа, 35); тренировки Wakesurf/Wakeskim — ул. Торфяная, 83",
        disciplines="Вейкборд–катер; Катер–доска длинная (Wakesurf); Катер–доска короткая (Wakeskim)",
        starts_at=datetime(2026, 8, 10, 9, 0, tzinfo=timezone.utc),
        ends_at=datetime(2026, 8, 17, 18, 0, tzinfo=timezone.utc),
        status=EventStatus.registration_open.value,
    )
    if event is None:
        event = Event(slug=slug, **payload)
        db.add(event)
    else:
        for key, value in payload.items():
            setattr(event, key, value)
    db.flush()
    return event


def clear_event_children(db, event_id: int) -> None:
    db.execute(delete(TrainingSlot).where(TrainingSlot.event_id == event_id))
    db.execute(delete(Official).where(Official.event_id == event_id))
    db.execute(delete(Participant).where(Participant.event_id == event_id))
    db.execute(delete(Document).where(Document.event_id == event_id))
    db.execute(delete(Category).where(Category.event_id == event_id))
    db.flush()


def _map_form_discipline(raw: str | None) -> str | None:
    if not raw:
        return None
    key = re.sub(r"\s+", " ", str(raw).strip().casefold())
    if key in FORM_DISCIPLINE_MAP:
        return FORM_DISCIPLINE_MAP[key]
    for alias, discipline in FORM_DISCIPLINE_MAP.items():
        if alias in key or key in alias:
            return discipline
    if "wakeskim" in key or "коротк" in key:
        return "Wakeskim"
    if "wakesurf" in key or "длинн" in key:
        return "Wakesurf"
    if "вейкборд" in key or "wakeboard" in key:
        return "Wakeboard (boat)"
    return None


def _infer_gender(category: str | None) -> str | None:
    if not category:
        return None
    low = category.casefold()
    if "женщин" in low or "девуш" in low or "девоч" in low:
        return "Ж"
    if "мужчин" in low or "юнош" in low or "мальчи" in low:
        return "М"
    return None


def _medical_from_cell(value) -> tuple[bool, str | None]:
    """Return (has_cert, url). Never log URL."""
    if value is None:
        return False, None
    text = str(value).strip()
    if not text or text.lower() in {"нет", "n/a", "-", "—"}:
        return False, None
    if text.startswith("http://") or text.startswith("https://"):
        return True, text[:1024]
    return True, None


def _get_or_create_category(
    db, event: Event, category_map: dict[tuple[str, str], Category], discipline: str, cat_title: str
) -> Category:
    key = (discipline, cat_title)
    if key not in category_map:
        code = f"{_slugify(discipline)}-{_slugify(cat_title)}"[:120]
        cat = Category(
            event_id=event.id,
            code=code,
            title=cat_title,
            discipline=discipline,
        )
        db.add(cat)
        db.flush()
        category_map[key] = cat
    return category_map[key]


def import_participants_from_form(db, event: Event, xlsx: Path) -> tuple[int, int, int]:
    """Primary import: Google Form answers (one row per discipline application).

    user_id is intentionally left NULL — one athlete may have several rows
    (multi-discipline) and UniqueConstraint(event_id, user_id) would break.
    """
    wb = load_workbook(xlsx, data_only=True, read_only=True)
    sheet = wb[wb.sheetnames[0]]
    category_map: dict[tuple[str, str], Category] = {}
    participants = 0
    with_medical = 0
    seen: set[tuple[int, str]] = set()
    headers: dict[str, int] | None = None

    for row_idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        values = list(row)
        if headers is None:
            maybe = _parse_headers(values)
            if "fio" in maybe and "discipline" in maybe:
                headers = maybe
            continue

        fio_idx = headers.get("fio")
        if fio_idx is None or fio_idx >= len(values) or not values[fio_idx]:
            continue
        fio = str(values[fio_idx]).strip()
        if not fio or fio.lower() in {"фио", "№"}:
            continue

        disc_raw = None
        if "discipline" in headers and headers["discipline"] < len(values):
            disc_raw = values[headers["discipline"]]
        discipline = _map_form_discipline(str(disc_raw) if disc_raw else None)
        if not discipline:
            continue

        cat_title = "Без категории"
        if "category" in headers and headers["category"] < len(values):
            raw = values[headers["category"]]
            if raw:
                cat_title = str(raw).strip()

        cat = _get_or_create_category(db, event, category_map, discipline, cat_title)
        dedupe_key = (cat.id, fio.casefold())
        if dedupe_key in seen:
            continue
        seen.add(dedupe_key)

        region = None
        if "region" in headers and headers["region"] < len(values) and values[headers["region"]]:
            region = str(values[headers["region"]]).strip()

        birth_year = None
        if "birth" in headers and headers["birth"] < len(values):
            birth_year = _birth_year(values[headers["birth"]])

        phone = None
        if "phone" in headers and headers["phone"] < len(values):
            phone = normalize_phone(values[headers["phone"]])

        has_med, med_url = False, None
        if "medical" in headers and headers["medical"] < len(values):
            has_med, med_url = _medical_from_cell(values[headers["medical"]])
            if has_med:
                with_medical += 1

        db.add(
            Participant(
                event_id=event.id,
                category_id=cat.id,
                full_name=fio,
                gender=_infer_gender(cat_title),
                birth_year=birth_year,
                region=region,
                phone=phone,
                has_medical_cert=has_med,
                medical_cert_url=med_url,
                status="registered",
                source_row=row_idx,
            )
        )
        participants += 1

    db.flush()
    return len(category_map), participants, with_medical


def import_participants(db, event: Event, xlsx: Path) -> tuple[int, int]:
    """Legacy category workbook (sheets per discipline)."""
    wb = load_workbook(xlsx, data_only=True, read_only=True)
    category_map: dict[tuple[str, str], Category] = {}
    participants = 0
    seen: set[tuple[int, str]] = set()

    for sheet_name, discipline in SHEET_DISCIPLINE.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        headers: dict[str, int] | None = None
        for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            values = list(row)
            if headers is None:
                maybe = _parse_headers(values)
                if "fio" in maybe:
                    headers = maybe
                continue
            if _is_section_row(values):
                continue
            fio_idx = headers.get("fio")
            if fio_idx is None or fio_idx >= len(values) or not values[fio_idx]:
                continue
            fio = str(values[fio_idx]).strip()
            if not fio or fio.lower() in {"фио", "№"}:
                continue
            if re.fullmatch(r"\d+", fio):
                continue

            cat_title = None
            if "category" in headers and headers["category"] < len(values):
                raw = values[headers["category"]]
                if raw:
                    cat_title = str(raw).strip()
            if not cat_title:
                cat_title = "Без категории"

            cat = _get_or_create_category(db, event, category_map, discipline, cat_title)

            dedupe_key = (cat.id, fio.casefold())
            if dedupe_key in seen:
                continue
            seen.add(dedupe_key)

            gender = None
            if "gender" in headers and headers["gender"] < len(values) and values[headers["gender"]]:
                gender = str(values[headers["gender"]]).strip()

            region = None
            if "region" in headers and headers["region"] < len(values) and values[headers["region"]]:
                region = str(values[headers["region"]]).strip()

            birth_year = None
            if "birth" in headers and headers["birth"] < len(values):
                birth_year = _birth_year(values[headers["birth"]])

            phone = None
            if "phone" in headers and headers["phone"] < len(values):
                phone = normalize_phone(values[headers["phone"]])

            db.add(
                Participant(
                    event_id=event.id,
                    category_id=cat.id,
                    full_name=fio,
                    gender=gender,
                    birth_year=birth_year,
                    region=region,
                    phone=phone,
                    status="registered",
                    source_row=row_idx,
                )
            )
            participants += 1

    db.flush()
    return len(category_map), participants


def load_phone_book_from_training(xlsx: Path) -> dict[str, str]:
    """Map normalized FIO -> phone from sheet «ЧП России»."""
    if not xlsx.exists():
        return {}
    wb = load_workbook(xlsx, data_only=True, read_only=True)
    sheet_name = next((n for n in wb.sheetnames if "росси" in n.casefold()), None)
    if sheet_name is None:
        return {}
    ws = wb[sheet_name]
    book: dict[str, str] = {}
    for idx, row in enumerate(ws.iter_rows(values_only=True)):
        if idx == 0:
            continue
        values = list(row)
        if len(values) < 5:
            continue
        fio = values[1]
        phone_raw = values[4]
        if not fio or not phone_raw:
            continue
        phone = normalize_phone(phone_raw)
        if not phone:
            continue
        for key in _fio_keys(str(fio)):
            book[key] = phone
    return book


def apply_phones_and_seed_users(db, event: Event, phone_book: dict[str, str]) -> tuple[int, int]:
    """Fill missing phones from training sheet + create one login user per unique phone."""
    attached = 0
    users_created = 0
    phone_owners: dict[str, str] = {}
    participants = db.scalars(select(Participant).where(Participant.event_id == event.id)).all()

    for part in participants:
        if not part.phone:
            phone = None
            for key in _fio_keys(part.full_name):
                phone = phone_book.get(key)
                if phone:
                    break
            if phone:
                part.phone = phone
                attached += 1
        if part.phone and part.phone not in phone_owners:
            phone_owners[part.phone] = part.full_name

    # Training sheet may list athletes who are not yet in form rows.
    for fio_key, phone in phone_book.items():
        if phone not in phone_owners:
            phone_owners[phone] = fio_key.title()

    owner_phone = "+79160117179"
    for phone, display_name in phone_owners.items():
        if phone == owner_phone:
            continue
        existing = db.scalar(select(User).where(User.phone == phone))
        if existing is not None:
            if not existing.display_name:
                existing.display_name = display_name
            continue
        email = f"p{phone.lstrip('+')}@participants.mywave.local"
        if db.scalar(select(User).where(User.email == email)):
            digest = hashlib.sha1(phone.encode("utf-8")).hexdigest()[:6]
            email = f"p{phone.lstrip('+')}.{digest}@participants.mywave.local"
        db.add(
            User(
                email=email,
                phone=phone,
                role="participant",
                requested_role="participant",
                status="active",
                display_name=display_name,
            )
        )
        users_created += 1
    db.flush()
    return attached, users_created


def _as_date(value) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _as_time(value) -> time | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.time().replace(microsecond=0)
    if isinstance(value, time):
        return value.replace(microsecond=0)
    text = str(value).strip()
    for fmt in ("%H:%M:%S", "%H:%M"):
        try:
            return datetime.strptime(text, fmt).time()
        except ValueError:
            continue
    return None


def _classify_slot_label(label: str | None) -> tuple[str, str | None]:
    if not label:
        return "open", None
    low = label.lower().strip()
    if "резерв" in low or "резвн" in low:
        return "reserve_day", label
    if "заправ" in low:
        return "refuel", label
    if "погруз" in low:
        return "loading", label
    return "blocked", label


def import_training_slots(db, event: Event, xlsx: Path) -> tuple[int, int]:
    """Import wakeboard/wakesurf schedule sheets. Skip 'ЧП России' (registration+PII)."""
    if not xlsx.exists():
        print(f"SKIP training file missing: {xlsx}")
        return 0, 0

    wb = load_workbook(xlsx, data_only=True)
    sheet_meta = {
        "Вейкборд": ("Wakeboard (boat)", "Оз. Кабан, ул. Хади Такташа, 35"),
        "Вейксерф": ("Wakesurf", "ул. Торфяная, 83"),
    }
    total = 0
    booked = 0
    sort_order = 0

    for sheet_name, (discipline, venue) in sheet_meta.items():
        if sheet_name not in wb.sheetnames:
            continue
        rows = list(wb[sheet_name].iter_rows(values_only=True))
        if len(rows) < 3:
            continue

        date_row = rows[1]
        date_left = _as_date(date_row[1] if len(date_row) > 1 else None)
        date_right = _as_date(date_row[4] if len(date_row) > 4 else None)
        if date_left is None:
            date_left = date(2026, 8, 11)
        if date_right is None:
            date_right = date(2026, 8, 12)

        for row in rows[2:]:
            vals = list(row)
            t_left = _as_time(vals[0] if vals else None)
            label_left = None if len(vals) < 2 or vals[1] is None else str(vals[1]).strip()
            t_right = _as_time(vals[3] if len(vals) > 3 else None)
            athlete = None if len(vals) < 5 or vals[4] is None else str(vals[4]).strip()
            note = None if len(vals) < 6 or vals[5] is None else str(vals[5]).strip()

            if t_left is not None:
                status, extra = _classify_slot_label(label_left)
                db.add(
                    TrainingSlot(
                        event_id=event.id,
                        discipline=discipline,
                        venue=venue,
                        slot_date=date_left,
                        slot_time=t_left,
                        status=status,
                        athlete_name=None,
                        notes=extra,
                        sort_order=sort_order,
                    )
                )
                sort_order += 1
                total += 1

            if t_right is not None:
                if athlete and athlete.lower() not in {"заправка", "погрузка"}:
                    status = "booked"
                    athlete_name = athlete
                    notes = note
                    booked += 1
                else:
                    status, extra = _classify_slot_label(athlete or note)
                    athlete_name = None
                    notes = extra or note
                db.add(
                    TrainingSlot(
                        event_id=event.id,
                        discipline=discipline,
                        venue=venue,
                        slot_date=date_right,
                        slot_time=t_right,
                        status=status,
                        athlete_name=athlete_name,
                        notes=notes,
                        sort_order=sort_order,
                    )
                )
                sort_order += 1
                total += 1

    db.flush()
    return total, booked


def import_documents(db, event: Event, docs: list[tuple[Path, str, str, str]]) -> int:
    dest_dir = ROOT / "data" / "documents" / event.slug
    dest_dir.mkdir(parents=True, exist_ok=True)
    count = 0
    for src, kind, language, title in docs:
        if not src.exists():
            print(f"SKIP missing doc: {src}")
            continue
        safe_name = re.sub(r"[^\w.\-()А-Яа-яЁё ]+", "_", src.name).strip()
        dest = dest_dir / safe_name
        shutil.copy2(src, dest)
        rel = f"{event.slug}/{safe_name}"
        db.add(
            Document(
                event_id=event.id,
                title=title,
                kind=kind,
                language=language,
                file_name=safe_name,
                relative_path=rel,
                description="Загружено seed-скриптом из локальных файлов организатора",
            )
        )
        count += 1
    db.flush()
    return count


def _email_for_official(full_name: str, role: str) -> str:
    base = re.sub(r"[^a-z0-9]+", ".", full_name.lower().replace("ё", "e"))
    # Latinize simple Cyrillic approx for unique local emails
    table = str.maketrans(
        {
            "а": "a",
            "б": "b",
            "в": "v",
            "г": "g",
            "д": "d",
            "е": "e",
            "ж": "zh",
            "з": "z",
            "и": "i",
            "й": "y",
            "к": "k",
            "л": "l",
            "м": "m",
            "н": "n",
            "о": "o",
            "п": "p",
            "р": "r",
            "с": "s",
            "т": "t",
            "у": "u",
            "ф": "f",
            "х": "h",
            "ц": "c",
            "ч": "ch",
            "ш": "sh",
            "щ": "sch",
            "ы": "y",
            "э": "e",
            "ю": "yu",
            "я": "ya",
            "ь": "",
            "ъ": "",
        }
    )
    slug = full_name.lower().translate(table)
    slug = re.sub(r"[^a-z0-9]+", ".", slug).strip(".")
    if not slug:
        slug = hashlib.sha1(full_name.encode("utf-8")).hexdigest()[:8]
    return f"{slug}.{role}@example.com"


def import_officials(db, event: Event) -> int:
    count = 0
    for sort_order, full_name, position, region, judge_category, role in OFFICIALS:
        user_id = None
        if role:
            email = _email_for_official(full_name, role)
            user = db.scalar(select(User).where(User.email == email))
            if user is None:
                user = User(
                    email=email,
                    role=role,
                    status="active",
                    requested_role=role,
                    display_name=full_name,
                )
                db.add(user)
                db.flush()
            else:
                user.role = role
                user.status = "active"
                user.requested_role = role
                user.display_name = full_name
                db.flush()
            user_id = user.id
        db.add(
            Official(
                event_id=event.id,
                sort_order=sort_order,
                full_name=full_name,
                position=position,
                region=region,
                judge_category=judge_category,
                notes="Импортировано из протокола №32-КС от 13.07.2026",
                user_id=user_id,
            )
        )
        count += 1
    db.flush()
    return count


def remove_demo_cup(db) -> None:
    demo = db.scalar(select(Event).where(Event.slug == "demo-cup"))
    if demo is None:
        return
    clear_event_children(db, demo.id)
    db.delete(demo)
    db.flush()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--registration", type=Path, default=DEFAULT_REG)
    parser.add_argument(
        "--legacy-registration",
        type=Path,
        default=DEFAULT_LEGACY_REG,
        help="Optional category workbook if form export is missing",
    )
    parser.add_argument("--reset-db", action="store_true", help="Delete sqlite file before seed")
    args = parser.parse_args()

    settings = get_settings()
    if args.reset_db and settings.database_url.startswith("sqlite:///"):
        for candidate in [
            API_DIR / "data" / "mywave_event.db",
            ROOT / "data" / "mywave_event.db",
            Path("data") / "mywave_event.db",
        ]:
            if candidate.exists():
                candidate.unlink()
                print(f"Removed {candidate}")

    get_settings.cache_clear()
    init_db()
    SessionLocal = get_session_factory()
    db = SessionLocal()
    try:
        ensure_organizer(db)
        remove_demo_cup(db)
        event = upsert_event(db)
        clear_event_children(db, event.id)

        medical_count = 0
        source = "form"
        if args.registration.exists():
            cats, parts, medical_count = import_participants_from_form(
                db, event, args.registration
            )
        elif args.legacy_registration.exists():
            source = "legacy"
            cats, parts = import_participants(db, event, args.legacy_registration)
            print(f"WARN: form export missing, used legacy categories: {args.legacy_registration}")
        else:
            raise FileNotFoundError(
                f"No registration file: {args.registration} or {args.legacy_registration}"
            )

        phone_book = load_phone_book_from_training(TRAINING_XLSX)
        phones_linked, users_from_phones = apply_phones_and_seed_users(db, event, phone_book)
        # Re-bind owner after participant phone users (collision safety).
        ensure_organizer(db)
        docs = import_documents(db, event, DEFAULT_DOCS)
        offs = import_officials(db, event)
        slots_total, slots_booked = import_training_slots(db, event, TRAINING_XLSX)
        db.commit()
        print(
            f"SEED OK event_id={event.id} slug={event.slug} source={source} "
            f"categories={cats} participants={parts} medical_flags={medical_count} "
            f"documents={docs} officials={offs} training_slots={slots_total} "
            f"booked={slots_booked} phones_linked={phones_linked} "
            f"phone_users={users_from_phones} phone_book_size={len(phone_book)}"
        )
        print("NOTE: medical URLs stored internally only; public API exposes has_medical_cert bool.")
        print("NOTE: sheet CHP Rossii used only for phone/FIO gap-fill (not for slots).")
        print("NOTE: OTP for seed users goes to mail_outbox / email; SMS later.")
        print("Judges can still use generated *@example.com emails via /auth/dev-login.")
        return 0
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


if __name__ == "__main__":
    raise SystemExit(main())
