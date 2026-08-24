"""Analyze source event files -> UTF-8 report (no PII dump to console)."""

from __future__ import annotations

import json
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

REPORT = Path(__file__).resolve().parents[2] / "data" / "seed_analysis.json"
REG = Path(r"c:\Users\X230\Downloads\Telegram Desktop\Регистрация_ЧР_ПР_2026_по_категориям.xlsx")
MAIL = Path(r"c:\Users\X230\Downloads\Реестр_рассылки_участникам_ЧР_2026.xlsx")
BULLETIN = Path(
    r"c:\Users\X230\Downloads\Telegram Desktop\Бюллетень_Первенство_Чемпионат_России_2026_Wakeboard_Wakesurf_Wakeskim (2).pdf"
)
RULES = [
    Path(r"c:\Users\X230\Downloads\2026-IWWF-Wakeboard-Boat-Rules-FINAL.pdf"),
    Path(r"c:\Users\X230\Downloads\2026-OFFICIAL-IWWF-WAKESURF-RULES-FINAL.pdf"),
]


def sheet_preview(path: Path, max_rows: int = 5) -> dict:
    wb = load_workbook(path, data_only=True, read_only=True)
    out: dict = {"file": path.name, "sheets": {}}
    for name in wb.sheetnames:
        ws = wb[name]
        headers = None
        sample = []
        rows = 0
        nonempty = 0
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            vals = [None if c is None else str(c).strip() for c in row]
            if all(v is None or v == "" for v in vals):
                continue
            nonempty += 1
            if headers is None:
                headers = vals
            elif len(sample) < max_rows:
                sample.append(vals[:20])
            rows += 1
        out["sheets"][name] = {
            "rows_with_content": nonempty,
            "headers": headers,
            "sample_rows": sample,
        }
    return out


def category_stats(path: Path) -> dict:
    wb = load_workbook(path, data_only=True, read_only=True)
    stats: dict = {}
    for name in wb.sheetnames:
        ws = wb[name]
        headers = None
        cats: Counter[str] = Counter()
        names = 0
        for row in ws.iter_rows(values_only=True):
            vals = list(row)
            if headers is None:
                headers = [str(c).strip() if c is not None else "" for c in vals]
                continue
            if all(c is None or str(c).strip() == "" for c in vals):
                continue
            names += 1
            # Heuristic: find category-like columns
            for idx, h in enumerate(headers):
                hl = h.lower()
                if any(k in hl for k in ("категор", "дисципл", "класс", "группа")):
                    if idx < len(vals) and vals[idx] is not None:
                        cats[str(vals[idx]).strip()] += 1
        stats[name] = {
            "headers": headers,
            "participants_approx": names,
            "category_counts": dict(cats.most_common(80)),
        }
    return stats


def pdf_info(path: Path) -> dict:
    if not path.exists():
        return {"file": str(path), "exists": False}
    reader = PdfReader(str(path))
    texts = []
    for page in reader.pages[:3]:
        texts.append((page.extract_text() or "")[:1500])
    return {
        "file": path.name,
        "exists": True,
        "pages": len(reader.pages),
        "text_sample": texts,
    }


def main() -> int:
    REPORT.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "registration": sheet_preview(REG) if REG.exists() else {"missing": str(REG)},
        "registration_stats": category_stats(REG) if REG.exists() else {},
        "mailing": sheet_preview(MAIL) if MAIL.exists() else {"missing": str(MAIL)},
        "bulletin": pdf_info(BULLETIN),
        "rules": [pdf_info(p) for p in RULES],
    }
    REPORT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"WROTE {REPORT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
