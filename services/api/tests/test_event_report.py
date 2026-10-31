"""EventReport AI assemble / Excel / role actions (ADR-0012)."""

from __future__ import annotations

import io

from openpyxl import load_workbook

from conftest import auth_header


def _seed_event_with_published_result(client, slug: str = "report-cup"):
    org = auth_header(client, f"{slug}-org@example.com", "organizer")
    part = auth_header(client, f"{slug}-athlete@example.com", "participant")
    judge = auth_header(client, f"{slug}-judge@example.com", "judge")

    created = client.post(
        "/api/v1/events",
        headers=org,
        json={
            "slug": slug,
            "title": "Report Cup Kazan",
            "status": "draft",
        },
    )
    assert created.status_code == 201, created.text
    event_id = created.json()["id"]
    client.patch(
        f"/api/v1/events/{event_id}/status",
        headers=org,
        json={"status": "registration_open"},
    )

    app = client.post(
        f"/api/v1/events/{event_id}/applications",
        headers=part,
        json={"category_id": None, "region": "Казань"},
    )
    assert app.status_code == 201, app.text
    participant_id = app.json()["id"]
    client.patch(
        f"/api/v1/events/{event_id}/applications/{participant_id}",
        headers=org,
        json={"status": "accepted"},
    )

    heat = client.post(
        f"/api/v1/events/{event_id}/heats",
        headers=org,
        json={"code": "F1", "title": "Final", "heat_number": 1},
    )
    assert heat.status_code == 201, heat.text
    heat_id = heat.json()["id"]
    fill = client.post(
        f"/api/v1/events/{event_id}/heats/{heat_id}/start-list/fill",
        headers=org,
        json={},
    )
    assert fill.status_code in (200, 201), fill.text
    items = fill.json()["items"]
    assert items
    entry_id = items[0]["id"]
    client.patch(
        f"/api/v1/events/{event_id}/heats/{heat_id}/start-list/{entry_id}/status",
        headers=org,
        json={"status": "dns"},
    )

    part2 = auth_header(client, f"{slug}-athlete2@example.com", "participant")
    app2 = client.post(
        f"/api/v1/events/{event_id}/applications",
        headers=part2,
        json={"category_id": None, "region": "Москва"},
    )
    pid2 = app2.json()["id"]
    client.patch(
        f"/api/v1/events/{event_id}/applications/{pid2}",
        headers=org,
        json={"status": "accepted"},
    )
    draft = client.post(
        f"/api/v1/events/{event_id}/results",
        headers=org,
        json={"participant_id": pid2, "score": 91.5, "place": 1, "heat_id": heat_id},
    )
    assert draft.status_code == 201, draft.text
    result_id = draft.json()["id"]
    client.patch(
        f"/api/v1/events/{event_id}/results/{result_id}/status",
        headers=org,
        json={"status": "verified"},
    )
    client.patch(
        f"/api/v1/events/{event_id}/results/{result_id}/status",
        headers=org,
        json={"status": "published"},
    )

    return org, judge, event_id, participant_id


def test_event_report_ai_assemble_publish_void(client):
    org, _judge, event_id, _pid = _seed_event_with_published_result(client, "ai-cup")

    missing = client.get(f"/api/v1/events/{event_id}/event-report", headers=org)
    assert missing.status_code == 404

    assembled = client.post(
        f"/api/v1/events/{event_id}/event-report/assemble",
        headers=org,
        json={},
    )
    assert assembled.status_code == 200, assembled.text
    body = assembled.json()
    assert body["status"] == "draft"
    assert body["source"] == "llm"
    assert body["facts"]["ai"]["provider"] in ("local", "openai")
    assert body["facts"]["published_results_count"] >= 1
    assert body["facts"]["dns_dnf_count"] >= 1
    assert body["facts"]["role_actions"]
    assert len(body["body_markdown"].strip()) > 40
    if body["facts"]["ai"]["provider"] == "local":
        assert "Призёры" in body["body_markdown"]
        assert "Передача действий" in body["body_markdown"]

    template = client.post(
        f"/api/v1/events/{event_id}/event-report/assemble",
        headers=org,
        json={"use_llm": False},
    )
    assert template.status_code == 200
    assert template.json()["source"] == "template"
    assert "Призёры" in template.json()["body_markdown"]

    client.post(f"/api/v1/events/{event_id}/event-report/assemble", headers=org, json={})
    published = client.post(f"/api/v1/events/{event_id}/event-report/publish", headers=org)
    assert published.status_code == 200
    assert published.json()["status"] == "published"

    conflict = client.post(
        f"/api/v1/events/{event_id}/event-report/assemble",
        headers=org,
        json={},
    )
    assert conflict.status_code == 409

    voided = client.post(f"/api/v1/events/{event_id}/event-report/void", headers=org)
    assert voided.status_code == 200
    assert voided.json()["status"] == "void"


def test_event_report_llm_disabled_when_flag_off(client, monkeypatch):
    from app.config import Settings, get_settings

    org, _judge, event_id, _pid = _seed_event_with_published_result(client, "flag-cup")

    monkeypatch.setenv("ENABLE_EVENT_REPORT_LLM", "0")
    get_settings.cache_clear()
    disabled = Settings()

    app = client.app
    app.dependency_overrides[get_settings] = lambda: disabled
    try:
        res = client.post(
            f"/api/v1/events/{event_id}/event-report/assemble",
            headers=org,
            json={"use_llm": True},
        )
        assert res.status_code == 400
        assert res.json()["error"]["code"] == "llm_disabled"
    finally:
        app.dependency_overrides.pop(get_settings, None)
        monkeypatch.delenv("ENABLE_EVENT_REPORT_LLM", raising=False)
        get_settings.cache_clear()


def test_event_report_excel_and_role_actions(client):
    org, _judge, event_id, _pid = _seed_event_with_published_result(client, "xlsx-cup")
    client.post(f"/api/v1/events/{event_id}/event-report/assemble", headers=org, json={})

    xlsx = client.get(f"/api/v1/events/{event_id}/event-report/excel", headers=org)
    assert xlsx.status_code == 200, xlsx.text
    assert "spreadsheetml" in xlsx.headers.get("content-type", "")
    wb = load_workbook(io.BytesIO(xlsx.content))
    assert "Summary" in wb.sheetnames
    assert "Podium" in wb.sheetnames
    assert "RoleActions" in wb.sheetnames
    assert "Narrative" in wb.sheetnames

    actions = client.get(f"/api/v1/events/{event_id}/role-actions", headers=org)
    assert actions.status_code == 200
    payload = actions.json()
    assert payload["total"] >= 10
    codes = {i["code"] for i in payload["items"]}
    assert "checkin_desk" in codes
    assert "event_report" in codes

    day = client.get(f"/api/v1/events/{event_id}/role-actions?phase=day", headers=org)
    assert day.status_code == 200
    assert all(i["phase"] == "day" for i in day.json()["items"])


def test_event_report_participant_cannot_assemble(client):
    org, _judge, event_id, _pid = _seed_event_with_published_result(client, "nosy-cup")
    part = auth_header(client, "er-nosy@example.com", "participant")
    res = client.post(
        f"/api/v1/events/{event_id}/event-report/assemble",
        headers=part,
        json={},
    )
    assert res.status_code == 403


def test_event_report_judge_cannot_publish(client):
    org, judge, event_id, _pid = _seed_event_with_published_result(client, "judge-cup")
    assemble = client.post(
        f"/api/v1/events/{event_id}/event-report/assemble",
        headers=judge,
        json={},
    )
    assert assemble.status_code == 403

    client.post(f"/api/v1/events/{event_id}/event-report/assemble", headers=org, json={})
    pub = client.post(f"/api/v1/events/{event_id}/event-report/publish", headers=judge)
    assert pub.status_code == 403
