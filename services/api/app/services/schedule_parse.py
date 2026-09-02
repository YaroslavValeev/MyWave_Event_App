"""Parse the «Расписание (3 дня)» sheet of the Kazan registration workbook."""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO

from openpyxl import load_workbook

DISCIPLINE_HINTS = (
    ("wakesurf", "Wakesurf"),
    ("длинн", "Wakesurf"),
    ("wakeskim", "Wakeskim"),
    ("коротк", "Wakeskim"),
    ("вейкборд", "Wakeboard (boat)"),
    ("wakeboard", "Wakeboard (boat)"),
)


@dataclass
class ParsedScheduleHeat:
    day_no: int
    day_date: datetime
    time_from: str
    discipline: str
    category_title: str
    round_key: str
    round_label: str
    heat_index: int
    athlete_count: int | None
    title: str


def _discipline(raw: str) -> str | None:
    key = raw.casefold()
    for alias, disc in DISCIPLINE_HINTS:
        if alias in key:
            return disc
    return None


def _round_key(raw: str) -> tuple[str, str]:
    low = (raw or "").casefold()
    if "lcq" in low:
        return "lcq", "LCQ"
    if "финал" in low:
        return "final", "Финал"
    return "qual", "Квалификация"


def _heat_index(raw: str) -> int:
    match = re.search(r"(\d+)", raw or "")
    return int(match.group(1)) if match else 1


def parse_schedule_workbook(payload: bytes) -> list[ParsedScheduleHeat]:
    wb = load_workbook(BytesIO(payload), data_only=True, read_only=True)
    sheet_name = next((n for n in wb.sheetnames if "распис" in n.casefold()), None)
    if sheet_name is None:
        wb.close()
        return []
    ws = wb[sheet_name]
    day_no = 0
    day_date: datetime | None = None
    rows: list[ParsedScheduleHeat] = []
    for values in ws.iter_rows(values_only=True):
        cells = ["" if c is None else str(c).strip() for c in values]
        first = cells[0] if cells else ""
        day_m = re.search(r"ДЕНЬ\s+(\d+)\s+[—\-]\s+(\d{2})\.(\d{2})\.(\d{4})", first)
        if day_m:
            day_no = int(day_m.group(1))
            day_date = datetime(
                int(day_m.group(4)),
                int(day_m.group(3)),
                int(day_m.group(2)),
            )
            continue
        if not day_date or not first:
            continue
        if first.casefold().startswith("время") or "заправка" in first.casefold() or "итого" in first.casefold():
            continue
        time_m = re.match(r"(\d{1,2}:\d{2})", first)
        if not time_m:
            continue
        discipline_raw = cells[1] if len(cells) > 1 else ""
        if "заправка" in discipline_raw.casefold() or "перерыв" in discipline_raw.casefold():
            continue
        disc = _discipline(discipline_raw)
        if not disc:
            continue
        cat_title = cells[2] if len(cells) > 2 else ""
        if not cat_title:
            continue
        round_key, round_label = _round_key(cells[3] if len(cells) > 3 else "")
        heat_index = _heat_index(cells[4] if len(cells) > 4 else "")
        count = None
        if len(cells) > 5 and cells[5].isdigit():
            count = int(cells[5])
        title = f"{disc} · {cat_title} · {round_label}"
        if heat_index > 1:
            title += f" · заезд {heat_index}"
        rows.append(
            ParsedScheduleHeat(
                day_no=day_no,
                day_date=day_date,
                time_from=time_m.group(1),
                discipline=disc,
                category_title=cat_title,
                round_key=round_key,
                round_label=round_label,
                heat_index=heat_index,
                athlete_count=count,
                title=title,
            )
        )
    wb.close()
    return rows
