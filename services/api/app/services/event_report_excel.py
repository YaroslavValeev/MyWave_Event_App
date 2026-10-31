"""Excel export of EventReport — derived view, not Source of Truth."""

from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font

from app.models.event_report import EventReport
from app.services.event_report_service import _parse_facts


def build_event_report_xlsx(*, event_slug: str, event_title: str, report: EventReport) -> bytes:
    facts = _parse_facts(report.facts_json)
    wb = Workbook()

    summary = wb.active
    summary.title = "Summary"
    summary.append(["Поле", "Значение"])
    for key, value in (
        ("event_slug", event_slug),
        ("event_title", event_title),
        ("report_status", report.status),
        ("report_source", report.source),
        ("report_title", report.title),
        ("assembled_at", facts.get("assembled_at")),
        ("ai_provider", (facts.get("ai") or {}).get("provider")),
        ("published_results_count", facts.get("published_results_count")),
        ("dns_dnf_count", facts.get("dns_dnf_count")),
        ("heats_count", facts.get("heats_count")),
        ("moments_status", facts.get("moments_status")),
        ("note", "Excel — выгрузка; SoT = API/БД Event App"),
    ):
        summary.append([key, value])
    for cell in summary[1]:
        cell.font = Font(bold=True)

    _sheet_rows(
        wb,
        "Podium",
        ["place", "full_name", "score", "participant_id", "heat_id", "category_id"],
        facts.get("podium") or [],
    )
    _sheet_rows(
        wb,
        "Results",
        ["place", "full_name", "score", "status", "participant_id"],
        facts.get("published_results") or [],
    )
    _sheet_rows(
        wb,
        "DNS_DNF",
        ["status", "full_name", "heat_code", "heat_id", "participant_id"],
        facts.get("dns_dnf") or [],
    )
    _sheet_rows(
        wb,
        "Officials",
        ["full_name", "position"],
        facts.get("officials") or [],
    )
    _sheet_rows(
        wb,
        "RoleActions",
        ["phase", "code", "role", "handoff_to", "physical_ru", "app_action"],
        facts.get("role_actions") or [],
    )

    narrative = wb.create_sheet("Narrative")
    narrative.append(["title", report.title])
    narrative.append(["body_markdown"])
    narrative.append([report.body_markdown or ""])
    narrative["A1"].font = Font(bold=True)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _sheet_rows(wb: Workbook, title: str, headers: list[str], rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([row.get(h) for h in headers])
