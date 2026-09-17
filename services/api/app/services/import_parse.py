"""Parse registration workbooks into normalized rows. No DB, no logging of PII."""

from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass, field
from datetime import datetime
from io import BytesIO

from openpyxl import load_workbook

from app.services.phone_utils import normalize_phone

HEADER_ALIASES = {
    "fio": {"фио", "ф.и.о.", "русск", "fullname"},
    "latin": {"латин", "latin", "лат."},
    "gender": {"пол"},
    "category": {"категория"},
    "region": {"регион"},
    "birth": {"дата рождения"},
    "phone": {"телефон", "phone", "мобил", "контактный"},
    "discipline": {"дисциплина"},
    "medical": {"медицин", "справка", "скан-копия"},
}

SHEET_DISCIPLINE = {
    "вейкборд-катер": "Wakeboard (boat)",
    "wakesurf (доска длинная)": "Wakesurf",
    "wakeskim (доска короткая)": "Wakeskim",
}

FORM_DISCIPLINE_MAP = {
    "вейкборд - катер": "Wakeboard (boat)",
    "вейкборд-катер": "Wakeboard (boat)",
    "катер - доска длинная (wakesurf)": "Wakesurf",
    "катер - доска короткая (wakeskim)": "Wakeskim",
}

SKIP_SHEETS = {"аналитика", "расписание (3 дня)"}


@dataclass
class ParsedRow:
    source_sheet: str
    source_row: int
    display_name: str | None
    latin_name: str | None
    phone_e164: str | None
    birth_year: int | None
    region: str | None
    discipline: str | None
    category_label: str | None
    has_medical: bool
    raw_safe: dict[str, str | int | bool | None] = field(default_factory=dict)


def content_sha256(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


def normalize_fio(value: str) -> str:
    return re.sub(r"\s+", " ", value.strip().casefold())


def map_discipline(raw: str | None) -> str | None:
    if not raw:
        return None
    key = re.sub(r"\s+", " ", str(raw).strip().casefold())
    if key in FORM_DISCIPLINE_MAP:
        return FORM_DISCIPLINE_MAP[key]
    for alias, discipline in FORM_DISCIPLINE_MAP.items():
        if alias in key or key in alias:
            return discipline
    if "wakeskim" in key or "коротк" in key:
        return "Wakeskim"
    if "wakesurf" in key or "длинн" in key:
        return "Wakesurf"
    if "вейкборд" in key or "wakeboard" in key:
        return "Wakeboard (boat)"
    return None


def _parse_headers(row: list) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for idx, cell in enumerate(row):
        if cell is None:
            continue
        label = str(cell).strip().lower()
        for key, aliases in HEADER_ALIASES.items():
            if label in aliases or any(a in label for a in aliases):
                mapping[key] = idx
    return mapping


def _birth_year(value) -> int | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.year
    text = str(value)
    match = re.search(r"(19|20)\d{2}", text)
    return int(match.group(0)) if match else None


def _cell(values: list, headers: dict[str, int], key: str) -> object | None:
    idx = headers.get(key)
    if idx is None or idx >= len(values):
        return None
    return values[idx]


def _safe_raw(headers: dict[str, int], values: list, *, has_medical: bool) -> dict:
    """Store labels without medical URLs or full phone digits."""
    out: dict[str, str | int | bool | None] = {"has_medical": has_medical}
    for key in ("fio", "latin", "category", "region", "discipline"):
        val = _cell(values, headers, key)
        if val is None or val == "":
            continue
        text = str(val).strip()
        out[key] = text[:255]
    birth = _birth_year(_cell(values, headers, "birth"))
    if birth:
        out["birth_year"] = birth
    return out


def _row_from_values(
    *,
    sheet: str,
    row_idx: int,
    values: list,
    headers: dict[str, int],
    discipline_hint: str | None,
) -> ParsedRow | None:
    fio_raw = _cell(values, headers, "fio")
    if not fio_raw:
        return None
    fio = str(fio_raw).strip()
    if not fio or fio.lower() in {"фио", "№"} or re.fullmatch(r"\d+", fio):
        return None
    disc_raw = _cell(values, headers, "discipline")
    discipline = map_discipline(str(disc_raw) if disc_raw else None) or discipline_hint
    cat_raw = _cell(values, headers, "category")
    category = str(cat_raw).strip() if cat_raw else None
    region_raw = _cell(values, headers, "region")
    latin_raw = _cell(values, headers, "latin")
    med_raw = _cell(values, headers, "medical")
    has_medical = False
    if med_raw is not None:
        text = str(med_raw).strip()
        has_medical = bool(text) and text.lower() not in {"нет", "n/a", "-", "—"}
    phone = normalize_phone(str(_cell(values, headers, "phone") or ""))
    return ParsedRow(
        source_sheet=sheet,
        source_row=row_idx,
        display_name=fio,
        latin_name=str(latin_raw).strip() if latin_raw else None,
        phone_e164=phone,
        birth_year=_birth_year(_cell(values, headers, "birth")),
        region=str(region_raw).strip() if region_raw else None,
        discipline=discipline,
        category_label=category,
        has_medical=has_medical,
        raw_safe=_safe_raw(headers, values, has_medical=has_medical),
    )


def parse_registration_workbook(payload: bytes) -> tuple[str, list[ParsedRow]]:
    """Return (source_kind, rows). kind is form_answers | category_roster."""
    wb = load_workbook(BytesIO(payload), data_only=True, read_only=True)
    rows: list[ParsedRow] = []
    first = wb.sheetnames[0]
    kind = "form_answers"

    # Form workbook: one sheet with discipline column.
    ws0 = wb[first]
    header_probe = None
    for values in ws0.iter_rows(values_only=True):
        header_probe = _parse_headers(list(values))
        break
    if header_probe and "fio" in header_probe and "discipline" in header_probe:
        headers: dict[str, int] | None = None
        for row_idx, row in enumerate(ws0.iter_rows(values_only=True), start=1):
            values = list(row)
            if headers is None:
                maybe = _parse_headers(values)
                if "fio" in maybe and "discipline" in maybe:
                    headers = maybe
                continue
            parsed = _row_from_values(
                sheet=first,
                row_idx=row_idx,
                values=values,
                headers=headers,
                discipline_hint=None,
            )
            if parsed:
                rows.append(parsed)
        wb.close()
        return kind, rows

    kind = "category_roster"
    for sheet_name in wb.sheetnames:
        if sheet_name.strip().casefold() in SKIP_SHEETS:
            continue
        hint = None
        folded = sheet_name.strip().casefold()
        for alias, disc in SHEET_DISCIPLINE.items():
            if alias in folded:
                hint = disc
                break
        ws = wb[sheet_name]
        headers = None
        for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            values = list(row)
            if headers is None:
                maybe = _parse_headers(values)
                if "fio" in maybe:
                    headers = maybe
                continue
            parsed = _row_from_values(
                sheet=sheet_name,
                row_idx=row_idx,
                values=values,
                headers=headers,
                discipline_hint=hint,
            )
            if parsed:
                rows.append(parsed)
    wb.close()
    return kind, rows


def dump_raw(parsed: ParsedRow) -> str:
    return json.dumps(parsed.raw_safe, ensure_ascii=False)
